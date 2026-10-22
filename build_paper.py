"""Build the revised research paper (.docx). Every number comes from results/results.json."""
import json
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

R = json.load(open("results/results.json"))
F = Path("figures")

f4 = lambda x: f"{float(x):.4f}"
f3 = lambda x: f"{float(x):.3f}"
pc = lambda x, n=1: f"{float(x) * 100:.{n}f}%"

T = {r["model"]: r for r in R["test_results"]}
CVM = {r["model"]: r for r in R["model_comparison_cv"]}
OURS, NN, LR, DT, LGB, TLGB, RF = ("Final model (ours)", "Neural Network / MLP (2014)", "Logistic Regression (2014)",
                                   "Decision Tree (2014)", "LightGBM", "Tuned LightGBM", "Random Forest")
o, nn = T[OURS], T[NN]
XGB, CAT, STK = "Tuned XGBoost", "CatBoost", "Stacked ensemble (LGB+HGB+LR)"
FN = R["final_model_name"]
DESC = {"Blend LightGBM+CatBoost": "an equal-weight blend of an Optuna-tuned LightGBM and a CatBoost model",
        "Blend XGBoost+CatBoost": "an equal-weight blend of an Optuna-tuned XGBoost and a CatBoost model",
        "Blend LightGBM+XGBoost+CatBoost": "an equal-weight blend of Optuna-tuned LightGBM and XGBoost and a CatBoost model",
        "Stacked ensemble (LGB+HGB+LR)": "a stacked ensemble of tuned LightGBM, histogram gradient boosting and logistic regression",
        "Tuned LightGBM": "an Optuna-tuned LightGBM", "Tuned XGBoost": "an Optuna-tuned XGBoost", "CatBoost": "a CatBoost model"}
fdesc = DESC[FN]
CANDS = {r["model"]: r for r in R["candidates_cv"]}
singles = [m for m in (TLGB, XGB, CAT) if m in T]
best_single = max(singles, key=lambda m: T[m]["pr_auc"])
best_2014 = max([LR, DT, NN], key=lambda m: T[m]["roc_auc"])
tuned_count = None   # set after Table 8
b14 = T[best_2014]
BS = R["bootstrap_ours_vs_mlp"]
GA = R["ga"]; SEL = {r["subset"]: r for r in R["feature_selection_test"]}
ga_row = SEL["GA-selected (ours)"]; all_row = SEL["All attributes"]
mi_row = next(v for k, v in SEL.items() if k.startswith("Mutual-info")); l1_row = SEL["L1-selected"]
KM = R["kmeans"]; STRAT = R["kmeans_strategy"]
LEAK = R["leakage"]; DS = R["dataset"]
REG = R["regularization"]; IMB = {r["strategy"]: r for r in R["imbalance_cv"]}
G50 = R["capture_top50pct"]; G20 = R["capture_top20pct"]
MV = R["missing_value_treatment_auc"]; PCA_ = R["pca"]
shap_top = list(R["shap_top15"].keys())
# Distinct drivers for the narrative: an 'unknown' indicator is the same signal as its parent attribute.
shap_concepts = list(dict.fromkeys(f[4:] if f.startswith("unk_") else f for f in shap_top))
TS = R["test_set"]; N_TE, POS_TE = TS["n"], TS["positives"]; NEG_TE = N_TE - POS_TE
def counts(r, at):
    # Confusion counts and accuracies exactly as computed in the notebook (no reconstruction).
    tp, fp, fn, tn = (int(r[f"{k}@{at}"]) for k in ("tp", "fp", "fn", "tn"))
    assert tp + fn == POS_TE and fp + tn == NEG_TE
    return dict(tp=tp, fp=fp, fn=fn, tn=tn, prec=tp / (tp + fp), rec=tp / POS_TE,
                acc=r[f"accuracy@{at}"], bal=r[f"bal_acc@{at}"])
c_nn05, c_nnt, c_o05, c_ot = counts(nn, "0.5"), counts(nn, "thr"), counts(o, "0.5"), counts(o, "thr")
always_no = TS["always_no_accuracy"]
# Which headline metrics favour our model over the 2014-style NN at its own tuned threshold (higher is better except Brier)
VS_TUNED = {"ROC-AUC": o["roc_auc"] - nn["roc_auc"], "PR-AUC": o["pr_auc"] - nn["pr_auc"], "recall": c_ot["rec"] - c_nnt["rec"],
            "precision": c_ot["prec"] - c_nnt["prec"], "F1": o["f1@thr"] - nn["f1@thr"], "balanced accuracy": c_ot["bal"] - c_nnt["bal"],
            "plain accuracy": c_ot["acc"] - c_nnt["acc"], "top-10% capture": o["top10_capture"] - nn["top10_capture"],
            "Brier score": nn["brier"] - o["brier"]}
behind = [k for k, v in VS_TUNED.items() if v < 0]
IMPROVED = [k for k, v in {"ranking quality": o["roc_auc"] - nn["roc_auc"], "recall": c_ot["rec"] - c_nn05["rec"],
                           "F1": o["f1@thr"] - nn["f1@0.5"], "balanced accuracy": c_ot["bal"] - c_nn05["bal"],
                           "calibration (Brier score)": nn["brier"] - o["brier"]}.items() if v > 0]
join = lambda xs: (", ".join(xs[:-1]) + " and " + xs[-1]) if len(xs) > 1 else (xs[0] if xs else "none")
def pfmt(p):   # bootstrap p-value; with B resamples the smallest resolvable two-sided p is 2/B
    return f"p < {2 / BS['n_resamples']:.3f}" if p == 0 else f"p = {p:.3f}"
NICE = {"contact": "contact channel", "month": "campaign month", "housing": "housing-loan status", "marital": "marital status",
        "prior_success": "previous campaign success", "poutcome": "previous outcome", "balance": "balance", "age": "age",
        "day": "day of month", "campaign": "number of calls in this campaign", "pdays": "days since previous contact",
        "unk_contact": "unrecorded contact channel", "job": "job", "education": "education", "loan": "personal loan",
        "days_since_contact": "days since previous contact", "never_contacted": "never-contacted flag",
        "previous": "number of previous contacts", "day_sin": "day-of-month (cyclical)", "month_sin": "month (cyclical)",
        "month_cos": "month (cyclical)", "balance_slog": "balance (signed log)"}
nice = lambda f: NICE.get(f, f.replace("_", " "))
LGB_IN_FINAL = "LightGBM" in FN

doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.27), Inches(11.69)
s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(0.95)
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.25
st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for lvl, sz in ((1, 14), (2, 12), (3, 11)):
    h = doc.styles[f"Heading {lvl}"]; h.font.name = "Times New Roman"; h.font.size = Pt(sz)
    h.font.bold = True; h.font.color.rgb = RGBColor(0, 0, 0)
    h.paragraph_format.space_before = Pt(12); h.paragraph_format.space_after = Pt(4)

def H(t, l=1): doc.add_heading(t, level=l)
def P(t, bold=False, italic=False, size=None, center=False):
    p = doc.add_paragraph(); r = p.add_run(t); r.bold, r.italic = bold, italic
    if size: r.font.size = Pt(size)
    if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p
def B(items):
    for it in items:
        p = doc.add_paragraph(it, style="List Bullet"); p.paragraph_format.space_after = Pt(3)
def table(headers, rows, caption, widths=None, size=9):
    c = doc.add_paragraph(); r = c.add_run(caption); r.bold = True; r.font.size = Pt(10)
    c.paragraph_format.keep_with_next = True; c.alignment = WD_ALIGN_PARAGRAPH.LEFT
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]; cell.text = ""; run = cell.paragraphs[0].add_run(str(h))
        run.bold = True; run.font.size = Pt(size)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        sh = OxmlElement("w:shd"); sh.set(qn("w:fill"), "E7ECF3"); cell._tc.get_or_add_tcPr().append(sh)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; run = cells[i].paragraphs[0].add_run(str(v)); run.font.size = Pt(size)
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
    for row in t.rows[:-1]:
        for cell in row.cells:
            for p_ in cell.paragraphs: p_.paragraph_format.keep_with_next = True
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); trPr.append(cs)
    if widths:
        t.autofit = False
        for i, w in enumerate(widths): t.columns[i].width = Inches(w)
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    doc.add_paragraph()
FIGN = [0]
def FIG(name, cap, w=6.2):
    FIGN[0] += 1
    doc.add_picture(str(F / name), width=Inches(w)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(); r = p.add_run(f"Figure {FIGN[0]}. "); r.bold = True; r.font.size = Pt(9.5)
    r2 = p.add_run(cap); r2.font.size = Pt(9.5); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return FIGN[0]
def pagebreak(): doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ======================================================== TITLE
P("A Leakage-Free Gradient-Boosting Ensemble with Genetic-Algorithm Feature Selection, "
  "Imbalance-Aware Thresholding and SHAP Explanations for Bank Telemarketing Success Prediction",
  bold=True, size=16, center=True)
P("Revisiting the Moro–Cortez–Rita (2014) bank-marketing model with modern machine learning",
  italic=True, size=11.5, center=True)
P("[Author One]¹, [Author Two]¹, [Author Three]¹", center=True)
P("¹ Department of Computer Engineering, Pimpri Chinchwad College of Engineering, Pune, India", size=9.5, center=True)
P("Corresponding author: [email]", size=9.5, center=True)

# ======================================================== ABSTRACT
H("Abstract")
P(f"Outbound telemarketing for term deposits succeeds in only about one call in eight, which makes accurate "
  f"pre-call targeting commercially valuable and technically difficult. A widely cited 2014 study by Moro, Cortez "
  f"and Rita compared logistic regression, decision trees, neural networks and support vector machines for this "
  f"task, but did not address class imbalance, reported only ranking metrics, and explained its model with "
  f"global sensitivity analysis. We revisit the problem on the public UCI Bank Marketing data "
  f"({DS['rows']:,} calls, {pc(DS['positive_rate'],2)} positive) with a fully leakage-free pipeline: call duration, "
  f"which is known only after the call, is excluded; hidden missing values are treated as informative; outliers are "
  f"capped; and all tuning, threshold selection and feature selection use the training split only. "
  f"We compare LightGBM, XGBoost and CatBoost, tuned with Optuna, together with stacked and blended ensembles, and "
  f"select the final model by cross-validation: {fdesc}, with a decision threshold chosen for F1. A genetic algorithm "
  f"(GA) selects the attribute subset. On an "
  f"untouched test set the proposed model reaches ROC-AUC {f4(o['roc_auc'])} and PR-AUC {f4(o['pr_auc'])}, against "
  f"{f4(nn['roc_auc'])} and {f4(nn['pr_auc'])} for a re-implemented 2014-style neural network (paired bootstrap {pfmt(BS['p_auc'])} for ROC-AUC). "
  f"Combined with the tuned threshold, it finds {pc(o['recall@thr'])} of subscribers, against {pc(nn['recall@0.5'])} for "
  f"the network at the default threshold (F1 {f3(o['f1@thr'])} vs {f3(nn['f1@0.5'])}); at matched tuned thresholds the "
  f"proposed model " + ("still leads" if o['f1@thr'] > nn['f1@thr'] else "does not lead") + f" (F1 {f3(o['f1@thr'])} vs {f3(nn['f1@thr'])}). Balanced accuracy rises from "
  f"{pc(c_nn05['bal'])} to {pc(c_ot['bal'])}; plain accuracy is reported but is misleading here, because predicting "
  f"'no' for every client already scores {pc(always_no)}. The GA keeps {GA['n_selected']} of {GA['n_total']} attributes at "
  f"test ROC-AUC {f4(ga_row['test_roc_auc'])}. K-means clustering on recency–frequency–monetary features identifies "
  f"{KM['k']} customer segments whose conversion rates range from "
  f"{pc(min(r['conversion'] for r in STRAT))} to {pc(max(r['conversion'] for r in STRAT))}, and TreeSHAP shows that "
  f"the {nice(shap_concepts[0])} and {nice(shap_concepts[1])} contribute most to the predictions.")
P("Keywords: bank telemarketing; class imbalance; gradient boosting; LightGBM; XGBoost; CatBoost; Optuna; genetic algorithm; "
  "feature selection; K-means; SHAP; data leakage; regularization", italic=True)

# ======================================================== 1 INTRO
H("1. Introduction")
P("Retail banks use outbound telemarketing to sell long-term deposits. Each call costs agent time, and most calls "
  "fail, so a model that ranks clients by their likelihood of subscribing directly reduces campaign cost. The UCI "
  "Bank Marketing dataset, released by Moro, Cortez and Rita, has become a standard benchmark for this problem [1, 2].")
P("In their 2014 study, Moro et al. [1] analysed 52,944 contacts from a Portuguese bank, reduced 150 candidate "
  "attributes to 22 with a semi-automatic selection procedure, and compared logistic regression (LR), decision trees "
  "(DT), a neural network (NN) and a support vector machine (SVM). Their best model, the NN, reached an AUC of 0.794 "
  "and an ALIFT of 0.672 in a rolling-window test, and they explained it through sensitivity analysis. Importantly, "
  "they deliberately excluded the duration of the current call because it is known only after the call ends.")
P("The study left several directions open, which this paper addresses:")
B(["Class imbalance. Only about 12% of calls succeed. The original work optimised AUC and did not report recall, "
   "precision, F1 or a decision threshold, so it is unknown how many real subscribers a deployed model would find.",
   "Model family. Gradient-boosted tree ensembles, which now dominate tabular prediction [7, 8, 15], were not compared.",
   "Feature selection. The 150→22 reduction relied on domain-guided, semi-automatic steps. A search-based method such "
   "as a genetic algorithm can explore attribute subsets systematically.",
   "Customer understanding. The original model scored clients individually but did not segment the customer base "
   "into groups for which different campaign strategies could be designed.",
   "Explainability. Global sensitivity analysis summarises the model as a whole but does not attribute an individual "
   "client's score to its inputs; SHAP values do [4]."])
P("Studies using the public release sometimes include call duration as an input, although the dataset's documentation "
  "advises discarding it for realistic prediction [2]; such models report inflated scores. We quantify that effect and "
  "exclude duration throughout.")
H("1.1 Contributions", 2)
B([f"A leakage-free, imbalance-aware gradient-boosting ensemble ({fdesc}) that, on the same data and split, improves PR-AUC by "
   f"{BS['delta_pr_auc']:+.4f} (95% CI {BS['pr_auc_ci95'][0]:+.4f} to {BS['pr_auc_ci95'][1]:+.4f}) and ROC-AUC by "
   f"{BS['delta_auc']:+.4f} (95% CI {BS['auc_ci95'][0]:+.4f} to {BS['auc_ci95'][1]:+.4f}) over a 2014-style neural "
   f"network. With an explicitly tuned threshold it finds {pc(o['recall@thr'])} of subscribers instead of "
   f"{pc(nn['recall@0.5'])} at the default threshold.",
   f"A genetic-algorithm feature selector over attribute groups, compared against all attributes, L1 selection and "
   f"mutual-information ranking.",
   "A systematic comparison of no penalty, L1, L2 and ElasticNet regularisation for the linear baseline.",
   f"An RFM-based K-means segmentation into {KM['k']} interpretable customer groups with targeted strategies.",
   "A complete data-preparation study (hidden missing values, outliers, encoding, correlation, PCA) and TreeSHAP "
   "explanations of the final model."])

# ======================================================== 2 RELATED
H("2. Related work")
P("Moro et al. [5] first applied the CRISP-DM methodology to this bank's campaigns, then extended it into the 2014 "
  "decision-support study [1]. The public dataset has since been used widely to test classifiers and resampling. "
  "SMOTE [6] and its variants are the most common imbalance treatments, although cost-sensitive weighting and "
  "threshold moving are often equally effective [12, 26]. Precision–recall analysis is preferred over ROC analysis when "
  "the positive class is rare [11]. Gradient boosting libraries such as XGBoost [7] and LightGBM [8] are the standard "
  "strong baselines for tabular data [15, 16]; CatBoost [29] adds ordered target statistics for categorical attributes, and "
  "Optuna [30] automates hyperparameter search. Stacking [27] combines heterogeneous learners through a meta-model. "
  "Genetic algorithms have long been used for wrapper feature selection [28]. SHAP [4] provides additive, "
  "game-theoretic attributions and has largely replaced perturbation-based sensitivity analysis for tree models. "
  "Kaufman et al. [10] formalise target leakage, of which post-outcome variables such as call duration are an example.")

# ======================================================== 3 DATA
pagebreak()
H("3. Dataset and preprocessing")
H("3.1 Data", 2)
P(f"We use bank-full.csv from the UCI repository [2]: {DS['rows']:,} records and {DS['cols']} attributes covering client "
  f"profile (age, job, marital status, education, credit default, balance, housing and personal loans), current "
  f"campaign (contact channel, day, month, number of calls) and previous campaigns (days since last contact, number of "
  f"previous contacts, previous outcome). The target is whether the client subscribed. Positives make up "
  f"{pc(DS['positive_rate'],2)}, an imbalance of {DS['imbalance_ratio']:.2f}:1. Monthly conversion varies from "
  f"{pc(R['month_conv_min_max'][0],2)} to {pc(R['month_conv_min_max'][1],2)}, so timing is a strong signal (Figure 1).")
FIG("lab1_distributions.png", "Distributions of age, balance and number of calls by outcome (top), and conversion rate "
    "by job, month and previous outcome (bottom). The dashed line marks the 11.7% base rate.")

H("3.2 Hidden missing values", 2)
unk = R["unknown"]
rows = []
for c in unk["unknown_count"]:
    rows.append([c, f"{int(unk['unknown_count'][c]):,}", f"{unk['unknown_pct'][c]:.1f}%",
                 f"{unk['conv_rate_if_unknown'][c]:.2f}%", f"{unk['conv_rate_if_known'][c]:.2f}%"])
table(["Attribute", "'unknown' count", "Share", "Conversion if unknown", "Conversion if known"], rows,
      "Table 1. Hidden missing values encoded as 'unknown'.")
P(f"The file contains no empty cells, but missing values are encoded as the string 'unknown' and as pdays = −1 for "
  f"clients never contacted before. Table 1 shows that 'unknown' clients convert at very different rates from the "
  f"rest, so the values are not missing at random. We therefore keep 'unknown' as its own category and add an "
  f"explicit indicator, rather than imputing the mode. In a check with logistic regression on the training rows, keeping the "
  f"category gave CV ROC-AUC {f4(MV['keep_unknown_as_category'])} against {f4(MV['mode_imputation'])} for mode imputation. "
  f"pdays = −1 is recoded to missing with a never_contacted flag, so the model is not told that these clients were "
  f"contacted one day ago.")

H("3.3 Outliers", 2)
rows = [[r["feature"], f"{r['lower']:.1f}", f"{r['upper']:.1f}", f"{r['n_outliers']:,}", f"{r['pct']:.2f}%"] for r in R["outliers"]]
table(["Feature", "Lower fence", "Upper fence", "Outliers", "Share"], rows, "Table 2. Outliers by the 1.5×IQR rule.")
P("The IQR rule flags many records (Table 2), but these are real clients, not errors. Deleting them would bias the "
  "sample, so extreme values are capped at the 1st and 99th percentiles of the training rows (winsorisation), and balance additionally "
  "receives a signed-log transform. The previous attribute is flagged heavily only because most clients have zero "
  "previous contacts.")

H("3.4 Encoding and feature engineering", 2)
P(f"Binary attributes are mapped to 0/1; education is also encoded ordinally; nominal attributes are one-hot encoded; "
  f"month and day are given cyclical sine–cosine encodings. Engineered features are total_contacts, prior_success, "
  f"days_since_contact, never_contacted, negative_balance, balance_slog, education_level and the four 'unknown' "
  f"indicators (unk_job, unk_education, unk_contact, unk_poutcome). The final matrix has "
  f"{R['n_features_encoded']} encoded columns. The duration attribute is removed from every model reported as ours.")

H("3.5 Correlation and feature relevance", 2)
tc = R["top_corr_with_target"]
P("Linear correlations with the target are weak for most attributes (Figure 2). The strongest is duration "
  f"(r = {tc.get('duration (post-call)', 0):.2f}), which is the post-call leakage variable. Among pre-call attributes the "
  "previous-campaign indicators are the strongest. Mutual information, which also captures non-linear dependence, "
  f"ranks {', '.join(R['mi_top15'][:5])} highest.")
FIG("lab1_correlation_heatmap.png", "Pearson correlation of numeric and engineered features with each other and with the target.", 5.6)

H("3.6 Principal component analysis", 2)
pl = list(PCA_["lr_auc"].items())
P(f"PCA [21] needs {PCA_['k90']} of {PCA_['n_features']} components for 90% of the variance and {PCA_['k95']} for 95%. The two "
  f"classes overlap heavily in the first two components (Figure 3). In cross-validation on the training rows, with PCA "
  f"fitted inside each fold, a logistic regression on the 90%-variance components scores ROC-AUC {f4(pl[1][1])} against {f4(pl[0][1])} on all features. PCA keeps high-variance "
  f"directions rather than predictive ones, so it is used for analysis and supervised selection is used for modelling.")
FIG("lab1_pca.png", "Left: cumulative explained variance. Right: projection on the first two principal components (orange = subscribed).")

# ======================================================== 4 METHODS
pagebreak()
H("4. Methodology")
H("4.1 Evaluation protocol", 2)
P("The data are split once into a stratified 80% training set and 20% test set. All model selection, "
  "hyperparameter tuning, decision-threshold selection, regularisation choice and GA search use five-fold (or, for the "
  "GA and tuning, three-fold) stratified cross-validation inside the training set only. Resampling such as SMOTE is "
  "applied inside each training fold through an imbalanced-learn pipeline [3]. The test set is used once, for the "
  "final numbers, which avoids optimistic selection bias [18, 19]. Outlier caps are also computed on training rows only. Metrics are ROC-AUC, PR-AUC (average precision), "
  "precision, recall, F1, Matthews correlation coefficient (MCC), accuracy, balanced accuracy, Brier score, and the share "
  "of all subscribers reached in the top 10%, 20% and 50% of scored clients. The difference between our model and the 2014-style network is tested with a 1,000-resample paired "
  "bootstrap on the test set [20].")
H("4.2 Baselines", 2)
P("The four 2014 model families are re-implemented on the same features and split: logistic regression with the "
  "best regularisation setting from Section 5.2, a depth-limited decision tree, a two-layer multilayer perceptron with early stopping, and a linear SVM "
  "with Platt calibration (a kernel SVM is impractical at this sample size). These are compared with random forest [9], "
  "histogram gradient boosting [25], LightGBM [8], XGBoost [7] and CatBoost [29].")
H("4.3 Regularisation", 2)
P("For logistic regression we compare no penalty, L1, L2 and ElasticNet (mixing ratio 0.5) across C ∈ {0.001, 0.01, "
  "0.1, 1, 10} [14, 17], recording training AUC, validation AUC, their gap (a measure of overfitting) and the number of non-zero "
  "coefficients.")
H("4.4 Class imbalance", 2)
P("Four strategies are compared for LightGBM: no treatment, balanced class weights, SMOTE inside each fold, and "
  "random undersampling. In addition, the decision threshold is moved from 0.5 to the value that maximises F1 on the "
  "training set's out-of-fold predictions.")
H("4.5 Proposed model", 2)
P(f"LightGBM and XGBoost hyperparameters (number of trees, learning rate, tree size, minimum leaf size, row and column "
  f"sampling, L1/L2 penalties) are tuned with Optuna's tree-structured Parzen estimator [30]: 40 trials each, three-fold "
  f"CV, PR-AUC objective. CatBoost [29] is used with fixed settings (1,000 iterations, depth 6, learning rate 0.04) and "
  f"receives the nominal attributes natively instead of one-hot columns. Seven candidates are then compared by five-fold "
  f"out-of-fold PR-AUC on the training set: the three single boosters, a stacked ensemble (LightGBM, histogram gradient "
  f"boosting and logistic regression with a logistic meta-learner), and three equal-weight probability blends of the "
  f"boosters. The candidate with the highest CV PR-AUC is selected as the final model: {fdesc}. Its decision threshold is "
  f"the F1-optimal value from training-set out-of-fold predictions ({f3(o['threshold'])}). The test set plays no part in "
  f"the selection.")
H("4.6 Genetic-algorithm feature selection", 2)
P(f"Each chromosome is a bit-string over the {GA['n_total']} attributes; all one-hot columns of an attribute are kept or "
  f"dropped together. Fitness is the three-fold CV ROC-AUC of a fast LightGBM (150 trees) minus 0.002 per selected "
  f"attribute, which rewards accuracy and penalises complexity. We use a population of 20, 15 generations, tournament "
  f"selection (size 3), uniform crossover (p = 0.8), bit-flip mutation (p = 1/L) and elitism of 2. The selected subset is "
  f"compared on the test set with all attributes, an L1-selected subset (logistic regression, C = 0.01) and a "
  f"mutual-information top-k subset of the same size, with one-hot columns mapped back to their parent attribute.")
H("4.7 RFM segmentation", 2)
P("Recency (days since the last previous contact, 999 if never), frequency (calls in this and previous campaigns) and "
  "monetary value (average balance) [24] are log-transformed; together with age they are standardised and clustered with "
  "K-means [22]. The number of clusters is chosen by the elbow (knee) rule, the k farthest from the straight line joining "
  "the first and last points of the normalised inertia curve, and the silhouette score is reported alongside [23]. The target is not used for clustering; "
  "it is used only afterwards to interpret the clusters.")
H("4.8 Explainability", 2)
P("Exact TreeSHAP values [4] are computed for the Optuna-tuned LightGBM" + (", a component of the final model," if LGB_IN_FINAL else "") + " on the test set with LightGBM's native implementation; "
  "contributions of one-hot columns are summed back to their parent attribute.")

# ======================================================== 5 RESULTS
pagebreak()
H("5. Results and discussion")
H("5.1 Summary of improvements", 2)
P("Table 3 brings the main result together: the 2014-style neural network, evaluated as in the reference study at the "
  "default 0.5 threshold and again with its own tuned threshold, against the final model of this work. All figures "
  "come from the same untouched test set.")
dlt = lambda a, b, pct=False: (f"{(b - a) * 100:+.1f} pts" if pct else f"{b - a:+.3f}")
summary_rows = [
    ["ROC-AUC", f3(nn["roc_auc"]), f3(nn["roc_auc"]), f3(o["roc_auc"]), dlt(nn["roc_auc"], o["roc_auc"])],
    ["PR-AUC", f3(nn["pr_auc"]), f3(nn["pr_auc"]), f3(o["pr_auc"]), dlt(nn["pr_auc"], o["pr_auc"])],
    ["Recall (subscribers found)", pc(c_nn05["rec"]), pc(c_nnt["rec"]), pc(c_ot["rec"]), dlt(c_nn05["rec"], c_ot["rec"], True)],
    ["Precision", pc(c_nn05["prec"]), pc(c_nnt["prec"]), pc(c_ot["prec"]), dlt(c_nn05["prec"], c_ot["prec"], True)],
    ["F1 score", f3(nn["f1@0.5"]), f3(nn["f1@thr"]), f3(o["f1@thr"]), dlt(nn["f1@0.5"], o["f1@thr"])],
    ["Balanced accuracy", pc(c_nn05["bal"]), pc(c_nnt["bal"]), pc(c_ot["bal"]), dlt(c_nn05["bal"], c_ot["bal"], True)],
    ["Accuracy*", pc(c_nn05["acc"]), pc(c_nnt["acc"]), pc(c_ot["acc"]), dlt(c_nn05["acc"], c_ot["acc"], True)],
    ["Subscribers in top 10% of calls", pc(nn["top10_capture"]), pc(nn["top10_capture"]), pc(o["top10_capture"]), dlt(nn["top10_capture"], o["top10_capture"], True)],
    ["Brier score (lower is better)", f4(nn["brier"]), f4(nn["brier"]), f4(o["brier"]), f"{o['brier'] - nn['brier']:+.4f}"],
    ["True subscribers identified", f"{c_nn05['tp']} / {POS_TE:,}", f"{c_nnt['tp']} / {POS_TE:,}", f"{c_ot['tp']} / {POS_TE:,}", f"{c_ot['tp'] - c_nn05['tp']:+d}"],
]
table(["Metric", "2014-style NN (t = 0.5)", f"2014-style NN (tuned t = {f3(nn['threshold'])})",
       f"Our model (tuned t = {f3(o['threshold'])})", "Change vs 2014 (t = 0.5)"], summary_rows,
      f"Table 3. Improvement summary on the held-out test set (n = {N_TE:,}, {POS_TE:,} subscribers). "
      f"*A model that predicts 'no' for every client scores {pc(always_no)} accuracy.", widths=[1.9, 1.05, 1.15, 1.15, 1.1], size=8.5)
P(f"{join(IMPROVED)[0].upper() + join(IMPROVED)[1:]} improve. The model identifies {c_ot['tp']} of the {POS_TE:,} subscribers in the "
  f"test set instead of {c_nn05['tp']}" + (", more than doubling recall," if c_ot['tp'] > 2 * c_nn05['tp'] else ",") + f" while F1 rises by {o['f1@thr'] - nn['f1@0.5']:.3f} and "
  f"balanced accuracy by {(c_ot['bal'] - c_nn05['bal']) * 100:.1f} percentage points. Even when the 2014-style network is "
  f"given its own tuned threshold, which was not part of the original method, our model remains ahead on every metric "
  f"except {join(behind)}. " + ("Precision falls because the tuned model calls more clients: it accepts more "
  f"unsuccessful calls in order to reach many more subscribers, the trade-off a campaign targeting buyers should make." if "precision" in behind else ""))
P(f"Plain accuracy needs care. Because only {pc(DS['positive_rate'])} of clients subscribe, a model that answers 'no' for "
  f"everyone reaches {pc(always_no)} accuracy while finding nobody. At the default threshold our model reaches "
  f"{pc(c_o05['acc'])} accuracy, against {pc(c_nn05['acc'])} for the network. At the tuned threshold its accuracy falls "
  f"to {pc(c_ot['acc'])}, because it deliberately calls more clients and accepts {c_ot['fp'] - c_o05['fp']} additional "
  f"false positives in exchange for {c_ot['tp'] - c_o05['tp']} additional subscribers. For a bank, a missed subscriber "
  f"costs far more than one extra call, so balanced accuracy, recall and F1 are the appropriate measures. The fall in "
  f"plain accuracy is the expected consequence of correcting for class imbalance, not a loss of model quality [11, 12].")

H("5.2 Regularisation", 2)
best_reg = max(REG, key=lambda r: r["cv_auc"]); none = [r for r in REG if r["penalty"] == "none"][0]
l1_small = min([r for r in REG if r["penalty"] == "l1" and r["cv_auc"] >= best_reg["cv_auc"] - 0.002], key=lambda r: r["nonzero_coefs"])
rows = [[r["penalty"], "—" if r["C"] != r["C"] else r["C"], f4(r["train_auc"]), f4(r["cv_auc"]), f4(r["gap"]), r["nonzero_coefs"]] for r in REG]
table(["Penalty", "C", "Train AUC", "CV AUC", "Gap", "Non-zero coefs"], rows, "Table 4. Regularisation of logistic regression (5-fold CV, training set).")
P(f"All settings show a train–validation gap below {max(r['gap'] for r in REG):.3f}, so the linear model does not "
  f"overfit; its limitation is bias. The best setting is {best_reg['penalty']} with C = {best_reg['C']} (CV ROC-AUC "
  f"{f4(best_reg['cv_auc'])}), essentially equal to no penalty ({f4(none['cv_auc'])}). Very strong regularisation lowers "
  f"AUC. L1 with C = {l1_small['C']} keeps only {l1_small['nonzero_coefs']} non-zero coefficients while staying within "
  f"0.002 AUC of the best, making it a useful embedded feature selector (Figure 4).")
FIG("lab3_regularization.png", "Validation AUC and number of non-zero coefficients against regularisation strength.")

H("5.3 Model comparison and class imbalance", 2)
FAM = {LR: "logistic regression", DT: "decision tree", NN: "neural network", "Linear SVM (2014)": "linear SVM"}
fam_best = FAM[max(FAM, key=lambda m: CVM[m]["roc_auc"])]; fam_worst = FAM[min(FAM, key=lambda m: CVM[m]["roc_auc"])]
rows = []
for m in [LR, DT, NN, "Linear SVM (2014)", RF, "Hist Gradient Boosting", LGB]:
    r = CVM.get(m)
    if r: rows.append([m, f4(r["roc_auc"]), f4(r["pr_auc"]), f3(r["recall"]), f3(r["f1"])])
table(["Model", "ROC-AUC", "PR-AUC", "Recall@0.5", "F1@0.5"], rows,
      "Table 5. Five-fold cross-validated comparison on the training set (no duration).")
P("Every model ranks clients reasonably well, but at the default 0.5 threshold all of them identify only a minority "
  "of subscribers (Table 5). Tree ensembles (random forest, histogram gradient boosting, LightGBM) lead on both ROC-AUC "
  "and PR-AUC. " + f"Among the 2014 families the {fam_best} is strongest and the {fam_worst} weakest. Moro et al. [1] also "
  "found the neural network best; in their rolling-window test logistic regression was weakest, while the decision tree "
  "was weakest in their modelling phase.")
rows = [[m + (" (selected)" if m == FN else ""), f4(r["roc_auc"]), f4(r["pr_auc"])] for m, r in CANDS.items()]
table(["Candidate", "CV ROC-AUC", "CV PR-AUC"], rows,
      "Table 6. Candidate final models, ranked by five-fold out-of-fold PR-AUC on the training set.", widths=[3.2, 1.2, 1.2])
cv_spread = max(r["roc_auc"] for r in CANDS.values()) - min(r["roc_auc"] for r in CANDS.values())
cand_sorted = list(CANDS)
P(f"All candidates fall within {cv_spread:.4f} ROC-AUC of each other (Table 6). The highest cross-validated PR-AUC comes "
  f"from {cand_sorted[0]} ({f4(CANDS[cand_sorted[0]]['pr_auc'])}), ahead of {cand_sorted[1]} "
  f"({f4(CANDS[cand_sorted[1]]['pr_auc'])}), but the margins are small: on this data the gradient-boosting family appears to be close to the ceiling that the available pre-call "
  f"attributes allow.")
rows = [[k, f4(v["roc_auc"]), f4(v["pr_auc"]), f3(v["precision"]), f3(v["recall"]), f3(v["f1"])] for k, v in IMB.items()]
table(["Strategy (LightGBM)", "ROC-AUC", "PR-AUC", "Precision@0.5", "Recall@0.5", "F1@0.5"], rows,
      "Table 7. Imbalance strategies (5-fold CV, training set).")
rmin = min(v["roc_auc"] for v in IMB.values()); rmax = max(v["roc_auc"] for v in IMB.values())
P(f"Imbalance treatment changes ROC-AUC by at most {rmax - rmin:.4f}, but changes recall at 0.5 from "
  f"{f3(min(v['recall'] for v in IMB.values()))} to {f3(max(v['recall'] for v in IMB.values()))} (Table 7). Resampling "
  f"mainly moves the operating point rather than improving the ranking. This confirms that reporting AUC alone, as in "
  f"the 2014 study, hides the quantity a campaign manager needs: how many buyers are found at the threshold used. "
  f"We therefore keep the unweighted model, whose probabilities stay calibrated, and move the threshold explicitly.")

H("5.4 Proposed model on the test set", 2)
rows = []
TBL8 = [m for m in [LR, DT, NN, RF, LGB, TLGB, XGB, CAT, STK] + ([FN] if FN not in (TLGB, XGB, CAT, STK) else []) if m in T]
for m in TBL8:
    if m not in T: continue
    r = T[m]
    rows.append([m + (" ★ ours" if m == FN else ""), f4(r["roc_auc"]), f4(r["pr_auc"]), f3(r["recall@0.5"]), f3(r["f1@0.5"]), f3(r["threshold"]),
                 f3(r["precision@thr"]), f3(r["recall@thr"]), f3(r["f1@thr"]), pc(r["top10_capture"])])
n_up = sum(T[m]["f1@thr"] > T[m]["f1@0.5"] for m in TBL8); tuned_txt = "all" if n_up == len(TBL8) else str(n_up)
table(["Model", "ROC-AUC", "PR-AUC", "Rec@0.5", "F1@0.5", "Thr*", "Prec@thr*", "Rec@thr*", "F1@thr*", "Top-10%"], rows,
      f"Table 8. Held-out test-set results (20%, n = {N_TE:,}). Thr* is the F1-optimal threshold chosen on training OOF predictions.",
      size=7.5)
P(f"On the untouched test set the selected model achieves ROC-AUC {f4(o['roc_auc'])} and PR-AUC {f4(o['pr_auc'])} "
  f"(Table 8). Against the 2014-style neural network the paired bootstrap ({BS['n_resamples']:,} resamples) gives ΔROC-AUC = "
  f"{BS['delta_auc']:+.4f} (95% CI {BS['auc_ci95'][0]:+.4f} to {BS['auc_ci95'][1]:+.4f}, {pfmt(BS['p_auc'])}) and "
  f"ΔPR-AUC = {BS['delta_pr_auc']:+.4f} (95% CI {BS['pr_auc_ci95'][0]:+.4f} to {BS['pr_auc_ci95'][1]:+.4f}, "
  f"{pfmt(BS['p_pr'])}).")
P(f"Compared with the best single booster on the test set ({best_single}: ROC-AUC {f4(T[best_single]['roc_auc'])}, "
  f"PR-AUC {f4(T[best_single]['pr_auc'])}), the selected model differs by {o['roc_auc'] - T[best_single]['roc_auc']:+.4f} "
  f"ROC-AUC and {o['pr_auc'] - T[best_single]['pr_auc']:+.4f} PR-AUC. " + ("Ensembling therefore adds stability rather than a "
  f"large improvement; " if FN not in singles else "") + f"most of the gain over the 2014 families comes from gradient boosting and tuning. We report the "
  f"model chosen by cross-validation even where another candidate happens to score slightly higher on the test set, "
  f"because choosing by test score would bias the estimate upward.")
P(f"The larger practical gain comes from the operating point. At the default 0.5 threshold the neural network finds "
  f"{pc(nn['recall@0.5'])} of subscribers (F1 {f3(nn['f1@0.5'])}). The proposed model at its tuned threshold of "
  f"{f3(o['threshold'])} finds {pc(o['recall@thr'])} with precision {pc(o['precision@thr'])} (F1 {f3(o['f1@thr'])}, MCC "
  f"{f3(o['mcc@thr'])}). Threshold tuning raises F1 for {tuned_txt} of the {len(TBL8)} models in Table 8, so for a fair comparison Table 8 also gives each baseline "
  f"its own tuned threshold: the network then reaches F1 {f3(nn['f1@thr'])} and recall {pc(nn['recall@thr'])}, " +
  ("still below our model." if o['f1@thr'] > nn['f1@thr'] else "matching or exceeding our model.") + f" Calling only the top 10% of clients ranked by our model reaches {pc(o['top10_capture'])} of all "
  f"subscribers ({pc(nn['top10_capture'])} for the network), and the top 20% reaches {pc(G20[OURS])} "
  f"({pc(G20[NN])}) (Figure 6). The 2014 study reported 79% of subscribers from the better half on its own, richer data "
  f"[1]; the top half of our ranking reaches {pc(G50[OURS])}. These figures are of the same order but are not directly "
  f"comparable with its 0.794 AUC, because the datasets and features differ.")
FIG("lab4_test_curves.png", "ROC and precision–recall curves on the test set, and the confusion matrix of the proposed model at its chosen threshold.", 6.5)
FIG("lab4_gains.png", "Cumulative gains: share of subscribers reached against share of clients called.", 4.3)

H("5.5 Genetic-algorithm feature selection", 2)
rows = [[k, v["n_attributes"], v["n_encoded_cols"], f4(v["test_roc_auc"]), f4(v["test_pr_auc"]), f"{v['fit_seconds']:.1f}"] for k, v in SEL.items()]
table(["Subset", "Attributes", "Encoded cols", "Test ROC-AUC", "Test PR-AUC", "Fit time (s)"], rows,
      "Table 9. Feature subsets evaluated with the tuned LightGBM on the test set.")
d_auc = ga_row["test_roc_auc"] - all_row["test_roc_auc"]
P(f"The GA ran for {len(GA['history'])} generations ({GA['evaluations']} unique subsets, "
  f"{GA['seconds']/60:.1f} minutes; Figure 7) and kept {GA['n_selected']} of {GA['n_total']} attributes, dropping "
  f"{', '.join(GA['dropped']) if GA['dropped'] else 'none'}. With these attributes the tuned LightGBM reaches test "
  f"ROC-AUC {f4(ga_row['test_roc_auc'])} versus {f4(all_row['test_roc_auc'])} with all attributes "
  f"(difference {d_auc:+.4f}), using {ga_row['n_encoded_cols']} instead of {all_row['n_encoded_cols']} encoded columns. "
  + f"PR-AUC, the more relevant metric for the rare positive class, is {f4(ga_row['test_pr_auc'])} against "
    f"{f4(all_row['test_pr_auc'])}. "
  + (f"The smaller model therefore costs {abs(d_auc):.4f} ROC-AUC and "
     f"{all_row['test_pr_auc'] - ga_row['test_pr_auc']:.4f} PR-AUC. It is the better choice when a simpler, faster model is "
     f"preferred; the full attribute set remains more accurate. "
     if ga_row['test_pr_auc'] < all_row['test_pr_auc'] else
     f"The GA model therefore matches the full model's PR-AUC with {ga_row['n_encoded_cols']} of {all_row['n_encoded_cols']} encoded columns. ")
  + (f"The L1 subset keeps {l1_row['n_attributes']} attributes and reaches ROC-AUC {f4(l1_row['test_roc_auc'])} and PR-AUC "
     f"{f4(l1_row['test_pr_auc'])}, so " + ("it ranks better than the GA subset but finds subscribers less precisely (lower PR-AUC). " if l1_row['test_roc_auc'] > ga_row['test_roc_auc'] and l1_row['test_pr_auc'] < ga_row['test_pr_auc'] else ("the GA subset is better on both metrics while being smaller. " if l1_row['test_roc_auc'] < ga_row['test_roc_auc'] and l1_row['test_pr_auc'] < ga_row['test_pr_auc'] else "it outperforms the GA subset, at the cost of more attributes. ")))
  + (f"Against a mutual-information subset of the same size (ROC-AUC {f4(mi_row['test_roc_auc'])}, PR-AUC "
     f"{f4(mi_row['test_pr_auc'])}), the GA subset is " + ("better on both metrics" if ga_row['test_roc_auc'] > mi_row['test_roc_auc'] and ga_row['test_pr_auc'] > mi_row['test_pr_auc'] else "comparable")
     + ", because it evaluates whole subsets and so accounts for redundancy and interactions between attributes, "
       "whereas mutual information scores each attribute on its own."))
FIG("lab7_ga_convergence.png", "GA convergence: best and mean fitness by generation.", 4.3)

H("5.6 Customer segmentation", 2)
rows = [[r["cluster"], r["segment"], pc(r["share"]), pc(r["conversion"]), r["strategy"]] for r in STRAT]
table(["Cluster", "Segment", "Share of base", "Conversion", "Recommended strategy"], rows,
      f"Table 10. K-means segments on RFM features (k = {KM['k']}, silhouette {KM['silhouette']:.3f} on a sample of {KM['silhouette_sample']:,} clients).", widths=[0.6, 1.5, 0.8, 0.8, 2.8])
KS = R["kmeans_selection"]
P(f"The elbow (knee) rule selects k = {KS['k_elbow']} (Figure 8). The silhouette score is highest at k = "
  f"{KS['k_best_silhouette']} ({KS['silhouette'][str(KS['k_best_silhouette'])]:.3f}) and is "
  f"{KS['silhouette'][str(KM['k'])]:.3f} at the chosen k" + ("; we keep the elbow choice because it separates more "
  f"actionable groups while the silhouette remains moderate." if KS['k_elbow'] != KS['k_best_silhouette'] else ", so both criteria agree.") + f" Every segment is named from its own profile. The highest-converting segment, "
  f"“{max(STRAT, key=lambda r: r['conversion'])['segment']}”, converts at {pc(max(r['conversion'] for r in STRAT))}, while the "
  f"weakest cold segment converts at {pc(min(r['conversion'] for r in STRAT))} (Table 10). Adding the cluster label as an "
  f"input to LightGBM changed CV ROC-AUC from {f4(R['cluster_as_feature']['without'])} to "
  f"{f4(R['cluster_as_feature']['with'])}, so segmentation adds little predictive power beyond the existing features. "
  f"Its value is strategic: it turns the scored base into groups that a marketing team can act on.")
FIG("lab5_elbow.png", "Elbow (inertia) and silhouette score for k = 2–10.", 5.4)
FIG("lab5_clusters.png", "Conversion rate per cluster and clusters in PCA space; cluster numbers match Table 10.", 5.8)

H("5.7 Explainability", 2)
SHAP_NOTE = {
    "contact": "Contact channel matters mainly through the unrecorded channel, whose clients convert at about 4%.",
    "month": "Campaign month reflects the strong seasonality of conversion (Figure 1).",
    "housing": "Clients with a housing loan are scored lower.",
    "prior_success": "A previous successful campaign is the strongest individual history signal.",
    "poutcome": "The previous campaign's outcome carries most of the contact-history signal.",
    "unk_contact": "An unrecorded contact channel lowers the score sharply, consistent with its 4% conversion rate.",
}
notes = " ".join(SHAP_NOTE[f] for f in shap_top[:4] if f in SHAP_NOTE and not (f == "contact" and "unk_contact" in shap_top[:4]))
P(f"TreeSHAP ranks {', '.join(nice(f) for f in shap_top[:6])} as the most influential attributes (Figure 10). {notes} "
  f"This agrees with the 2014 finding that call context and timing matter, although the macro-economic indicators "
  f"(e.g. Euribor rate) and agent attributes that ranked highest in [1] are not present in the public release and could "
  f"not be tested. Some information is encoded more than once: each 'unknown' indicator duplicates the 'unknown' level "
  f"of its one-hot attribute, month also appears as its sine–cosine encoding, and balance as its signed logarithm. "
  f"TreeSHAP splits credit among such encodings (for example unk_contact and contact), so they should be read together. SHAP values are attributions of the model's output and do not imply causal "
  f"effects.")
FIG("shap_summary.png", "TreeSHAP: mean absolute contribution per attribute (left) and per-client contributions (right; red = high, blue = low feature value, grey = categorical attribute).", 6.4)

H("5.8 Leakage through call duration", 2)
P(f"On the test set, duration alone gives ROC-AUC {f4(LEAK['duration_alone_auc_test'])}, and adding it to the tuned "
  f"LightGBM raises ROC-AUC from {f4(LEAK['without_duration_auc_test'])} to {f4(LEAK['with_duration_auc_test'])}. Studies "
  f"that include duration therefore report much higher scores, but such models cannot select whom to call. Moro et al. "
  f"excluded duration for this reason [1], and we follow them.")

H("5.9 Comparison with the 2014 study", 2)
table(["Aspect", "Moro et al. (2014)", "This work"], [
    ["Data", "52,944 calls, 150→22 attributes incl. macro-economic", f"{DS['rows']:,} calls, public 17-attribute release"],
    ["Models", "LR, DT, NN, SVM", "Same families + RF, HGB, LightGBM, XGBoost, CatBoost, stacking and blending (Optuna-tuned)"],
    ["Imbalance", "Not addressed", "Class weights, SMOTE, undersampling, threshold tuning"],
    ["Metrics", "AUC, ALIFT", "AUC, PR-AUC, precision, recall, F1, MCC, accuracy, balanced accuracy, Brier, top-k capture"],
    ["Feature selection", "Semi-automatic, domain-guided", "Genetic algorithm, L1, mutual information"],
    ["Segmentation", "None", "RFM + K-means with strategies"],
    ["Explainability", "Sensitivity analysis + rule tree", "TreeSHAP (global and per-client)"],
    ["Call duration", "Excluded", "Excluded; leakage quantified"]],
    "Table 11. Positioning against the reference study.", widths=[1.4, 2.4, 2.6])

H("5.10 Limitations", 2)
B(["The public dataset lacks the macro-economic and agent attributes of the original, so absolute scores cannot be "
   "compared directly with the reported 0.794 AUC.",
   "The data cover one Portuguese bank in 2008–2010; results may not transfer to other banks or periods.",
   "A single random 80/20 split is used for the final test; a chronological split would be more conservative given "
   "the strong seasonality.",
   "The GA fitness uses a fast LightGBM and a fixed complexity penalty; other settings could select different subsets.",
   "SHAP attributions describe the model, not causal effects of interventions [13]."])

# ======================================================== 6 CONCLUSION
H("6. Conclusion")
P(f"We revisited the 2014 bank-telemarketing model with a leakage-free, imbalance-aware pipeline. Among Optuna-tuned "
  f"LightGBM, XGBoost, CatBoost and their ensembles, cross-validation selects {fdesc}. With an F1-optimised threshold it "
  f"gives test ROC-AUC {f4(o['roc_auc'])} and PR-AUC {f4(o['pr_auc'])}, significantly above a 2014-style network, and finds "
  f"{pc(o['recall@thr'])} of subscribers instead of {pc(nn['recall@0.5'])} at the default threshold. The boosting "
  f"candidates lie close together, which suggests the public attributes, rather than the algorithm, now limit accuracy. "
  f"Imbalance treatment changes recall far more than AUC, so recall and PR-AUC should be reported alongside AUC. A "
  f"genetic algorithm reduces the attribute set to {GA['n_selected']} of {GA['n_total']} at ROC-AUC "
  f"{f4(ga_row['test_roc_auc'])}, K-means on RFM features yields {KM['k']} actionable segments, and TreeSHAP identifies "
  f"the {nice(shap_concepts[0])} and {nice(shap_concepts[1])} as the main drivers. Future work should apply the pipeline to the "
  f"bank-additional dataset, which includes the macro-economic indicators, use chronological validation, and add a "
  f"cost-sensitive threshold based on real call costs and deposit margins.")

H("Acknowledgements")
P("We thank S. Moro, P. Cortez and P. Rita for making the dataset publicly available through the UCI Machine Learning "
  "Repository.")
P("Data and code availability. The dataset is available at https://doi.org/10.24432/C5K306. The notebook that "
  "produces every number, table and figure in this paper is available at [repository URL].")
P("Declaration of generative AI assistance. [Adapt to your institution's policy.] An AI assistant was used for code "
  "scaffolding and drafting. All results were produced by the accompanying code and verified by the authors.", italic=True, size=9.5)

# ======================================================== REFERENCES
pagebreak()
H("References")
refs = [
 "S. Moro, P. Cortez, P. Rita, A data-driven approach to predict the success of bank telemarketing, Decision Support Systems 62 (2014) 22–31. https://doi.org/10.1016/j.dss.2014.03.001",
 "S. Moro, P. Cortez, P. Rita, Bank Marketing, UCI Machine Learning Repository, 2014. https://doi.org/10.24432/C5K306",
 "G. Lemaître, F. Nogueira, C.K. Aridas, Imbalanced-learn: a Python toolbox to tackle the curse of imbalanced datasets in machine learning, Journal of Machine Learning Research 18 (17) (2017) 1–5.",
 "S.M. Lundberg, S.-I. Lee, A unified approach to interpreting model predictions, Advances in Neural Information Processing Systems 30 (2017) 4765–4774.",
 "S. Moro, R. Laureano, P. Cortez, Using data mining for bank direct marketing: an application of the CRISP-DM methodology, Proc. European Simulation and Modelling Conference (2011) 117–121.",
 "N.V. Chawla, K.W. Bowyer, L.O. Hall, W.P. Kegelmeyer, SMOTE: synthetic minority over-sampling technique, Journal of Artificial Intelligence Research 16 (2002) 321–357.",
 "T. Chen, C. Guestrin, XGBoost: a scalable tree boosting system, Proc. 22nd ACM SIGKDD (2016) 785–794.",
 "G. Ke et al., LightGBM: a highly efficient gradient boosting decision tree, Advances in Neural Information Processing Systems 30 (2017) 3146–3154.",
 "L. Breiman, Random forests, Machine Learning 45 (1) (2001) 5–32.",
 "S. Kaufman, S. Rosset, C. Perlich, O. Stitelman, Leakage in data mining: formulation, detection, and avoidance, ACM Transactions on Knowledge Discovery from Data 6 (4) (2012) 1–21.",
 "T. Saito, M. Rehmsmeier, The precision-recall plot is more informative than the ROC plot when evaluating binary classifiers on imbalanced datasets, PLoS ONE 10 (3) (2015) e0118432.",
 "H. He, E.A. Garcia, Learning from imbalanced data, IEEE Transactions on Knowledge and Data Engineering 21 (9) (2009) 1263–1284.",
 "C. Rudin, Stop explaining black box machine learning models for high stakes decisions and use interpretable models instead, Nature Machine Intelligence 1 (5) (2019) 206–215.",
 "R. Tibshirani, Regression shrinkage and selection via the lasso, Journal of the Royal Statistical Society Series B 58 (1) (1996) 267–288.",
 "L. Grinsztajn, E. Oyallon, G. Varoquaux, Why do tree-based models still outperform deep learning on typical tabular data?, Advances in Neural Information Processing Systems 35 (2022) 507–520.",
 "R. Shwartz-Ziv, A. Armon, Tabular data: deep learning is not all you need, Information Fusion 81 (2022) 84–90.",
 "H. Zou, T. Hastie, Regularization and variable selection via the elastic net, Journal of the Royal Statistical Society Series B 67 (2) (2005) 301–320.",
 "S. Varma, R. Simon, Bias in error estimation when using cross-validation for model selection, BMC Bioinformatics 7 (2006) 91.",
 "G.C. Cawley, N.L.C. Talbot, On over-fitting in model selection and subsequent selection bias in performance evaluation, Journal of Machine Learning Research 11 (2010) 2079–2107.",
 "B. Efron, R.J. Tibshirani, An Introduction to the Bootstrap, Chapman & Hall, New York, 1993.",
 "I.T. Jolliffe, J. Cadima, Principal component analysis: a review and recent developments, Philosophical Transactions of the Royal Society A 374 (2016) 20150202.",
 "J. MacQueen, Some methods for classification and analysis of multivariate observations, Proc. 5th Berkeley Symposium on Mathematical Statistics and Probability 1 (1967) 281–297.",
 "P.J. Rousseeuw, Silhouettes: a graphical aid to the interpretation and validation of cluster analysis, Journal of Computational and Applied Mathematics 20 (1987) 53–65.",
 "A.M. Hughes, Strategic Database Marketing, Probus, Chicago, 1994.",
 "J.H. Friedman, Greedy function approximation: a gradient boosting machine, Annals of Statistics 29 (5) (2001) 1189–1232.",
 "C. Elkan, The foundations of cost-sensitive learning, Proc. 17th International Joint Conference on Artificial Intelligence (2001) 973–978.",
 "D.H. Wolpert, Stacked generalization, Neural Networks 5 (2) (1992) 241–259.",
 "J. Yang, V. Honavar, Feature subset selection using a genetic algorithm, IEEE Intelligent Systems 13 (2) (1998) 44–49.",
 "L. Prokhorenkova, G. Gusev, A. Vorobev, A.V. Dorogush, A. Gulin, CatBoost: unbiased boosting with categorical features, Advances in Neural Information Processing Systems 31 (2018) 6638–6648.",
 "T. Akiba, S. Sano, T. Yanase, T. Ohta, M. Koyama, Optuna: a next-generation hyperparameter optimization framework, Proc. 25th ACM SIGKDD (2019) 2623–2631.",
]
for i, r in enumerate(refs, 1):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.35); p.paragraph_format.first_line_indent = Inches(-0.35)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT; run = p.add_run(f"[{i}] {r}"); run.font.size = Pt(9.5)

out = "Research_Paper_Bank_Telemarketing_v2.docx"; doc.save(out); print("saved", out)
